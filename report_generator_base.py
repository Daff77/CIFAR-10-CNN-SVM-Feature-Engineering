"""
Script untuk meng-generate 8 Laporan Word (.docx) akademik lengkap
untuk proyek CIFAR-10 CNN + SVM Feature Engineering.
Setiap laporan bersumber 100% dari notebook final, artefak aktual, dan plot riil.
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Warna tema akademik UIN Sunan Ampel (Emerald Dark Green & Neutral Slate)
COLOR_PRIMARY_HEX = "2E5339"      # Hijau UIN
COLOR_SECONDARY_HEX = "1C3D29"    # Hijau Gelap
COLOR_BG_ALT_HEX = "F4F7F5"       # Abu-abu kehijauan lembut
COLOR_BORDER_HEX = "C5D0C8"       # Garis tabel halus
COLOR_CALLOUT_HEX = "EBF2EC"      # Kotak informasi

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, color=COLOR_BORDER_HEX, sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_base_document():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height
    
    # Atur default style Times New Roman
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)
    
    return doc

def add_cover(doc, pertemuan_title, subjudul, notebook_name):
    # Paragraph 1: Judul Laporan
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(4)
    run1 = p1.add_run(f"LAPORAN PRAKTIKUM PERTEMUAN 7\n{pertemuan_title.upper()}")
    run1.bold = True
    run1.font.name = "Times New Roman"
    run1.font.size = Pt(14)
    run1.font.color.rgb = RGBColor(46, 83, 57)
    
    # Paragraph 2: Subjudul / Modul
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(14)
    run2 = p2.add_run(f"{subjudul}\n(Notebook: {notebook_name})")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(12)
    run2.bold = True
    
    # Paragraph 3: Mata Kuliah & Dosen
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(8)
    p3.paragraph_format.space_after = Pt(24)
    run3 = p3.add_run("Mata Kuliah: Big Data Analytic\nDosen Pengampu: Bayu Adi Nugroho, Ph.D")
    run3.font.name = "Times New Roman"
    run3.font.size = Pt(12)
    
    # Paragraph 4: Logo UIN Sunan Ampel
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(10)
    p_logo.paragraph_format.space_after = Pt(28)
    if os.path.exists("uin_logo.png"):
        p_logo.add_run().add_picture("uin_logo.png", width=Inches(2.2))
    
    # Paragraph 5: Identitas Mahasiswa
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(14)
    p4.paragraph_format.space_after = Pt(4)
    run4 = p4.add_run("Disusun Oleh:")
    run4.font.name = "Times New Roman"
    run4.font.size = Pt(12)
    run4.bold = True
    
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(0)
    p5.paragraph_format.space_after = Pt(32)
    run5 = p5.add_run("Daffa Danendra Fairuzza – 09020624026")
    run5.font.name = "Times New Roman"
    run5.font.size = Pt(12)
    run5.bold = True
    
    # Paragraph 6: Institusi
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(20)
    p6.paragraph_format.space_after = Pt(0)
    run6 = p6.add_run(
        "PROGRAM STUDI SISTEM INFORMASI\n"
        "FAKULTAS SAINS DAN TEKNOLOGI\n"
        "UNIVERSITAS ISLAM NEGERI SUNAN AMPEL SURABAYA\n"
        "2026"
    )
    run6.font.name = "Times New Roman"
    run6.font.size = Pt(12)
    run6.bold = True
    
    doc.add_page_break()

def add_h1(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(46, 83, 57)
    return p

def add_h2(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11)
        r_b.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.italic = italic
    return p

def add_callout(doc, text, title="CATATAN PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.27)
    cell = table.cell(0, 0)
    set_cell_background(cell, COLOR_CALLOUT_HEX)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # border kiri tebal hijau
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_PRIMARY_HEX}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r_t = p.add_run(f"💡 {title}: ")
    r_t.bold = True
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(10.5)
    r_t.font.color.rgb = RGBColor(46, 83, 57)
    
    r_c = p.add_run(text)
    r_c.font.name = "Times New Roman"
    r_c.font.size = Pt(10.5)
    
    # spasi kosong setelah tabel
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11)
        r_b.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    return p

def add_table_data(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        set_cell_background(hdr_cells[i], COLOR_PRIMARY_HEX)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = COLOR_BG_ALT_HEX if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            # Angka/nilai di center, deskripsi di left
            if c_idx == 0 and len(headers) > 2 and len(str(val)) < 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif any(k in headers[c_idx].lower() for k in ['loss', 'accuracy', 'akurasi', 'precision', 'recall', 'f1', 'support', 'epoch', 'dimensi', 'shape']):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                
    # Atur lebar kolom jika ada
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table

def add_figure(doc, img_path, caption_text, width=Inches(5.6)):
    if not os.path.exists(img_path):
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    p_img.add_run().add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(9.5)
    r_cap.italic = True
    r_cap.font.color.rgb = RGBColor(60, 60, 60)

print("Base document generator module ready.")
